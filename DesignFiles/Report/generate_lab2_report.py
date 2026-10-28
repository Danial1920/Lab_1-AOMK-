from pathlib import Path
from shutil import copyfile
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'DesignFiles/Report/EcoLab2_Report.docx'
doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.7)
sec.left_margin=Cm(2);sec.right_margin=Cm(2)
sec.header_distance=Cm(.8);sec.footer_distance=Cm(.8)
styles=doc.styles
for name in ['Normal','Body Text','Caption']:
 st=styles[name];st.font.name='Arial';st.font.size=Pt(11)
 st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=1.12
for name,size in [('Title',25),('Heading 1',16),('Heading 2',12)]:
 st=styles[name];st.font.name='Arial';st.font.size=Pt(size)
 st.font.color.rgb=RGBColor.from_string('000000')
 st.font.bold=True;st.paragraph_format.space_after=Pt(10)
 st.paragraph_format.space_before=Pt(7 if name!='Title' else 0)
 st.paragraph_format.keep_with_next=True
for border in list(styles['Title']._element.xpath('.//w:pBdr')):
 border.getparent().remove(border)
styles['Caption'].font.size=Pt(9);styles['Caption'].font.color.rgb=RGBColor.from_string('4D5962')
styles['Caption'].paragraph_format.space_after=Pt(8)
styles['Caption'].paragraph_format.keep_with_next=False
header=sec.header.paragraphs[0];header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=header.add_run('ACOM  •  Лабораторная работа 2');r.font.name='Arial';r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('000000')
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

def p(text='',style=None):
 par=doc.add_paragraph(text,style);par.paragraph_format.widow_control=True;return par

def title(text):doc.add_heading(text,level=1)
def sub(text):doc.add_heading(text,level=2)
def page():doc.add_page_break()
def bullet(text):p(text,'List Bullet')
def code(text):
 par=p();par.paragraph_format.space_after=Pt(8);par.paragraph_format.line_spacing=1.05
 par.paragraph_format.keep_together=True
 r=par.add_run(text);r.font.name='Courier New';r.font.size=Pt(8.7)
 sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F2F5F7');par._p.get_or_add_pPr().append(sh)
 return par

def table(headers,rows,widths=None,size=9.5):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for i,h in enumerate(headers):t.rows[0].cells[i].text=h
 for row in rows:
  cells=t.add_row().cells
  for i,s in enumerate(row):cells[i].text=str(s)
 if widths:
  for row in t.rows:
   for c,w in zip(row.cells,widths):c.width=Cm(w)
 for j,row in enumerate(t.rows):
  pr=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(pr)
  for c in row.cells:
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcpr=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
   for side,val in [('top','75'),('bottom','75'),('left','95'),('right','95')]:
    m=OxmlElement('w:'+side);m.set(qn('w:w'),val);m.set(qn('w:type'),'dxa');mar.append(m)
   tcpr.append(mar)
   if j==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E6E6E6');tcpr.append(sh)
   elif j%2==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F7F7F7');tcpr.append(sh)
   for par in c.paragraphs:
    par.paragraph_format.space_after=Pt(1);par.paragraph_format.space_before=Pt(1)
    par.paragraph_format.line_spacing=1.03
    for r in par.runs:
     r.font.name='Arial';r.font.size=Pt(size)
     if j==0:r.font.bold=True;r.font.color.rgb=RGBColor(0,0,0)
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 p().paragraph_format.space_after=Pt(2)
 return t

def figure(name,width,caption):
 par=p();par.alignment=WD_ALIGN_PARAGRAPH.CENTER;par.paragraph_format.space_after=Pt(2)
 par.add_run().add_picture(str(ROOT/'DesignFiles/Report/figures'/name),width=Cm(width))
 cap=p(caption,'Caption');cap.alignment=WD_ALIGN_PARAGRAPH.CENTER

# The typography and cover follow the existing report; the text reflects the ABI revision.
p().paragraph_format.space_after=Pt(56)
par=p('Лабораторная работа 2','Title');par.alignment=WD_ALIGN_PARAGRAPH.CENTER
par=p('Включение и агрегирование\nкомпонентов ACOM');par.alignment=WD_ALIGN_PARAGRAPH.CENTER
for r in par.runs:r.font.size=Pt(18);r.font.bold=True
p().paragraph_format.space_after=Pt(22)
par=p('Eco.Lab1 и Eco.Lab2\nКалькуляторные компоненты A–E');par.alignment=WD_ALIGN_PARAGRAPH.CENTER
for r in par.runs:r.font.size=Pt(14)
p().paragraph_format.space_after=Pt(78)
par=p('Выполнил\nАнвари Д.А\n24ТИДИ');par.alignment=WD_ALIGN_PARAGRAPH.CENTER
for r in par.runs:r.font.size=Pt(14)
p().paragraph_format.space_after=Pt(45)
par=p('Нижний Новгород, 2026');par.alignment=WD_ALIGN_PARAGRAPH.CENTER

page();title('1 Постановка задачи и критерии')
p('Лабораторная работа расширяет компонент полиномиальной аппроксимации из ЛР1 калькуляторными интерфейсами. Реализованы включение и агрегирование ACOM-компонентов, внешний Eco.Lab2 и получение интерфейсов через QueryInterface. Проверка выполнена на macOS arm64 с официальным Eco Framework; исходники Framework не изменялись.')
p('Основание — «Задание_2.pdf». Для варианта на 9–10 баллов требуется внешний Eco.Lab2, агрегирование существующего Eco.Lab1 и демонстрация A, B, C, D, E, Eco.Lab1. Приведённые ниже результаты подтверждают проверенные свойства реализации; итоговую оценку определяет преподаватель.')
table(['Требование','Реализация'],[
 ('Калькулятор в прежнем компоненте','Lab1 агрегирует B для X и включает E для Y'),
 ('Включение и агрегирование A–E','D включает B, E агрегирует B; A/C заменяют друг друга по одному CID'),
 ('Внешнее агрегирование Lab1','Lab2 → Lab1 → B; общий controlling unknown'),
 ('Переходы QueryInterface','Матрицы всех поддерживаемых интерфейсов, включая X → Lab2'),
 ('Реальное взаимодействие','Horner вызывает Addition у A/C и Multiplication у D'),
 ('Отдельный клиент и серверы','Готовые архивы компонентов и один неизменяемый объект клиента')],widths=[5.1,11.9])
sub('Материалы и совместимость')
p('Контракты проверены по InsideACOM и репозиторию Zzzpize/EcoLab1, commit 57d07171c5c92b492ac3f1fd48b93ce2155dd67c. Применены официальные целочисленные X/Y и CID A–E. Клиент собирается по неизменённым учебным заголовкам; собственное безопасное деление вынесено в отдельный интерфейс.')
p('Функциональность ЛР1 сохранена: polyfit для пяти типов, нормальные уравнения/Гаусс и QR/Givens, FitVector, IEcoPolynomial и прежние CID. Численное ядро не менялось. Новые калькуляторные операции принимают int16, а не double; типы численного API polyfit остаются прежними.')

page();title('2 Компонентная архитектура')
figure('lab2_architecture.png',16.8,'Рисунок 1. Владение объектами и пути вызовов в Lab2')
p('Сплошные стрелки означают включение, пунктирные — агрегирование. B1, B2 и B3 — разные экземпляры одной фабрики. Все зависимости Eco-режима создаются через IEcoInterfaceBus1::QueryComponent. При отсутствии E Lab1 использует D; при отсутствии B компонент D включает A/C. Ошибка выделения памяти не запускает резервный путь.')
table(['Вызов','Реальный исполнитель'],[
 ('Lab2.QueryInterface(X)','B1.X из цепочки Lab2 → Lab1 → B1'),
 ('Lab2.QueryInterface(Y)','Фасад Lab1.Y → включённый E.Y'),
 ('Horner: произведение и сумма','D.Y.Multiplication и A/C.X.Addition'),
 ('D.QueryInterface(X)','Фасад D.X → включённый B3.X'),
 ('E.QueryInterface(X)','Агрегированный B2.X')],widths=[6.5,10.5],size=9.5)
p('C — независимая реализация X/Y, заменяющая A с тем же CID. В один runtime устанавливается A или C. Это проверяет заменяемость компонентов: остальные серверы и объект клиента при замене не перекомпилируются.')

page();title('3 Публичные контракты и идентификаторы')
p('X и Y содержат ровно пять слотов vtable: QueryInterface, AddRef, Release и две арифметические операции. Аргументы — int16; Addition/Multiplication возвращают int32, Subtraction/Division — int16. Division округляет к нулю. Subtraction явно приводит результат по модулю 65536.')
table(['Интерфейс','IID'],[
 ('IEcoCalculatorX','93221116-2248-4742-AE06-82819447843D'),
 ('IEcoCalculatorY','BD6414C2-9096-423E-A90C-04D77AFD1CAD'),
 ('IEcoCalculatorSafe','D3B3E2F1-8D21-4A6B-9A40-440000000054'),
 ('IEcoLab1','D3B3E2F1-8D21-4A6B-9A40-440000000052'),
 ('IEcoLab2 с Horner','D3B3E2F1-8D21-4A6B-9A40-440000000055')],widths=[5,12],size=9)
table(['Компонент','CID'],[
 ('A и C','4828F655-2E45-40E7-8121-EBD220DC360E'),
 ('B','AE202E54-3CE5-4550-8996-03BD70C62565'),
 ('D','3A8E4467-7E82-475C-B4A3-719ED8397E61'),
 ('E','872FEF1D-E331-4B87-AD44-D1E7C232C2F0'),
 ('Lab1 normal','D3B3E2F1-8D21-4A6B-9A40-440000000045'),
 ('Lab1 QR','D3B3E2F1-8D21-4A6B-9A40-440000000046'),
 ('Lab2','D3B3E2F1-8D21-4A6B-9A40-440000000065')],widths=[5,12],size=9)
sub('Безопасное деление и Horner')
p('IEcoCalculatorSafe::TryDivision возвращает status и результат через int16*. При b = 0 выдаётся 1101, при −32768 / −1 — 1102; выход обнуляется. Обычный Division в этих случаях возвращает 0. Метод Safe не добавлен в официальную таблицу Y.')
p('Horner принимает коэффициенты от старшей степени и возвращает status/int32 out. Перед передачей аргументов калькуляторам промежуточный аккумулятор и произведение проверяются на диапазон int16. Последнее сложение может дать int32, например 65534. Ошибка диапазона возвращает 1102 и нулевой выход. NULL, нулевая длина и переполнение размера отклоняются.')
p('Добавление Horner изменило IID Lab2 с …0053 на …0055. Заголовки и IDL согласованы; генерация межпроцессного маршалинга не входит в работу.')

page();title('4 Управление ссылками и временем жизни')
p('При включении владелец хранит обычную дочернюю ссылку. Дочерний объект сохраняет собственный IEcoUnknown. Фасад владельца вызывает методы дочернего интерфейса и сохраняет identity владельца.')
p('Агрегирующая фабрика принимает внешний unknown только вместе с IID_IEcoUnknown. Она возвращает приватный неделегирующий unknown, управляющий внутренним временем жизни. Публичные QI/AddRef/Release делегируют controlling unknown внешнего объекта.')
figure('lab2_lifetime.png',16.2,'Рисунок 2. Публичные ссылки и приватное владение агрегатом')
p('Lab2 создаёт Lab1 с outer = Lab2. Lab1 создаёт B с тем же controlling unknown. Поэтому B.X.QueryInterface(IEcoUnknown) возвращает identity Lab2, а AddRef/Release изменяют внешний счётчик. Независимый клиент освобождает исходную ссылку Lab2, продолжает вызывать B.X, получает Lab2 обратно через QI и завершает все ссылки.')
p('Последний внешний Release освобождает приватный unknown Lab1 и включённые A/C и D. Lab1 освобождает private B и включённый E; E — private B своего экземпляра. Каждый объект освобождает дочерние ссылки раньше system/allocator. Внешний код не приводит чужие интерфейсы к concrete-структурам и не вызывает чужие destroy-функции.')
p('Повторный Release уже уничтоженного объекта недопустим: проверки касаются баланса владеемых ссылок и отсутствия двойного освобождения, а не обращения к освобождённой памяти.')

page();title('5 Методика и результаты тестирования')
p('Результаты ниже получены 28 октября 2026. Host использует явно обозначенный тестовый allocator. Eco-клиент работает с настоящими Framework и Interface Bus; фабрики и QueryInterface не подменяются. Наблюдатель allocator считает выделения и позволяет вызвать детерминированный отказ Alloc.')
table(['Проверка','Фактический результат'],[
 ('Host','180 условий, 0 ошибок'),
 ('Eco с A / с C','1062 / 1110 условий, 0 ошибок'),
 ('Allocator в каждом Eco-варианте','267 выделений, 267 освобождений, live = 0'),
 ('Независимый клиент Debug и Release','A: 523 / 0; C: 559 / 0'),
 ('Dynamic Debug/Release × A/C','Все 4 запуска: A 523 / 0, C 559 / 0'),
 ('Xcode Static Debug и Release','Компонент и UnitTest собраны; запуск: 1062 / 0'),
 ('C89, libc, ASan/UBSan','Проверки пройдены, диагностик sanitizer нет'),
 ('Отрицательные мутации и сборочные зависимости','Ошибочная операция и abort() обнаружены'),
 ('Валидатор численных эталонов','8 проверок пройдены')],widths=[8,9],size=9.5)
p('Счётчики отражают количество условий, а не процент покрытия. Тесты возвращают ненулевой код при ошибке или неполном плане. Полные логи сохранены в DesignFiles/Review/Verification_2026-10-28.')
sub('Проверенные сценарии')
p('Для семи уникальных CID в каждом A/C-варианте выполнены матрицы QI, identity, счётчики ссылок и реальные операции. Отдельно проверены восемь фабрик, семь групп отказов Alloc, четыре случая отсутствующих зависимостей и два резервных пути. Агрегированные интерфейсы проверены после раннего освобождения корня.')
p('Публичный клиент проверяет полином Horner в восьми точках по независимому арифметическому эталону, границы int16/int32, NULL и ошибочные размеры. Polyfit проверяется через Lab1, QR и Lab2; существующие численные тесты ЛР1 сохранены.')
p('Мутации выполняются в временной копии проекта: неправильная реальная арифметика заставляет UnitTest завершиться ошибкой; добавленный abort() обнаруживается через nm. Изменение заголовков/флагов запускает пересборку, неверный путь Framework отклоняется.')

page();title('6 Сборка и независимый клиент')
p('Среда проверки — macOS arm64, Xcode и распакованный официальный Eco.Framework. Команды выполняются из корня репозитория последовательно. Путь Framework должен указывать на существующий каталог.')
code('export ECO_FRAMEWORK="$HOME/Downloads/Eco.Framework"\nmake ECO_FRAMEWORK="$ECO_FRAMEWORK" check\nmake ECO_FRAMEWORK="$ECO_FRAMEWORK" eco-test-variants\nCONFIGURATION=Release make ECO_FRAMEWORK="$ECO_FRAMEWORK" component-clients\nmake ECO_FRAMEWORK="$ECO_FRAMEWORK" validation-build')
p('check запускает Host, C89, проверку libc, ASan/UBSan для Host и Eco в обоих вариантах, валидатор эталонов и отдельный клиент Debug. eco-test-variants отдельно запускает обычную Eco-интеграцию для A/C.')
sub('Архивы и общий клиент')
p('build_component_clients.py создаёт восемь отдельных архивов: A, B, C, D, E, Lab1, Lab1QR и Lab2. Компиляция Lab2Client.c использует ExternalContracts, SharedFiles и официальные заголовки Framework; закрытый HeaderFiles и C-файлы серверов исключены. Один Lab2Client.o линкуется с семью архивами: сначала A, затем C. SHA256 объекта после обоих запусков совпадает.')
code('ECO_EXPECT_CALCULATOR=A BuildFiles/Components/Debug/EcoLab2ClientA\nECO_EXPECT_CALCULATOR=C BuildFiles/Components/Debug/EcoLab2ClientC')
sub('Xcode Static')
p('Схема Static компонента оставлена общим архивом ради совместимости проекта ЛР1. Отдельные архивы поставки проверяются предыдущим скриптом. Оба Xcode-проекта собираются в Debug/Release; после сборки UnitTest нужно запустить.')
code('for configuration in Debug Release; do\n  xcodebuild -project AssemblyFiles/Mac/Xcode_v123/EcoLab1.xcodeproj \\\n    -scheme Static -configuration "$configuration" -jobs 1 \\\n    -derivedDataPath "BuildFiles/XcodeStandardABI/Component$configuration" \\\n    ECO_FRAMEWORK="$ECO_FRAMEWORK" CODE_SIGNING_ALLOWED=NO build\n  xcodebuild -project AssemblyFiles/Mac/Xcode_v123/EcoLab1UnitTest.xcodeproj \\\n    -scheme Static -configuration "$configuration" -jobs 1 \\\n    -derivedDataPath "BuildFiles/XcodeStandardABI/Unit$configuration" \\\n    ECO_FRAMEWORK="$ECO_FRAMEWORK" CODE_SIGNING_ALLOWED=NO build\n  "BuildFiles/XcodeStandardABI/Unit$configuration/Build/Products/$configuration/Static"\ndone')
p('Нельзя запускать параллельные сборки в один derivedDataPath: это блокирует build.db. Сборки текущего прогона использовали отдельные каталоги для проекта и конфигурации.')

page();title('7 Dynamic runtime и ограничения')
code('for configuration in Debug Release; do\n  for variant in A C; do\n    ECO_FRAMEWORK="$ECO_FRAMEWORK" CONFIGURATION="$configuration" \\\n      CALCULATOR_VARIANT="$variant" \\\n      DesignFiles/Report/build_dynamic_runtime_macos.sh || break 2\n  done\ndone')
p('На каждый запуск создаётся семь CID-named dylib: A либо C, B, D, E, Lab1 normal, Lab1 QR и Lab2. A/C имеют одно имя файла и размещаются в разных runtime-каталогах. Отдельный Dynamic-клиент использует те же публичные тесты, но линкует только Framework и не регистрирует лабораторные фабрики.')
p('Официальный загрузчик получает библиотеки через ECO_FRAMEWORK_RT. DYLD trace подтверждает загрузку всех семи. Дополнительный процесс с пустым runtime обязан завершиться ошибкой. Для каждой dylib проверены nm -u и otool -L. Логи: BuildFiles/DynamicRuntimeLogs/<Debug|Release>/<A|C>.')
sub('Численные эталоны ЛР1')
p('Сохранённые прогоны от 25 октября 2026 сравнили 36 массивов коэффициентов SciPy/NumPy и 18 массивов MATLAB R2026b для девяти наборов и обоих алгоритмов. При текущей правке внешние MATLAB/SciPy не перезапускались; регрессионные Host/Eco-тесты выполнены заново. Повторные команды: make PYTHON=/opt/anaconda3/bin/python reference и matlab-reference.')
sub('Ограничения')
p('Проверена macOS arm64. Бинарные сборки Windows/Linux/EcoOS не выполнялись; отсутствие официального ECO_TOOLCHAIN обозначается SKIPPED. В архиве нет Eco.StdIO.C89 arm64, поэтому стандартный вывод используется только консольными клиентами. Компоненты не содержат StdIOShim и не вызывают libc; платформенная зависимость dylib от libSystem проверяется отдельно от собственных неопределённых символов.')
p('На macOS arm64 long double имеет точность double. Непредставимые точно long отклоняются численным API с 1005. Horner ограничен целочисленными аргументами учебных калькуляторов; полный диапазон произвольных полиномов не заявляется.')
sub('Вывод')
p('Реализация демонстрирует требуемые механизмы включения и агрегирования, вложенный controlling unknown и заменяемость A/C. Клиент проверен отдельно от реализации серверов, операции вызывают реальные дочерние интерфейсы, а отрицательные сценарии дают ошибку. Архитектурное сравнение не ограничивается именами папок: проверены ABI, владение ссылками, загрузка и наблюдаемое поведение.')
p('Источники: Задание_2.pdf; peerf-eco/Eco.Education, 001.InsideACOM; Zzzpize/EcoLab1 (commit указан в разделе 1); официальные заголовки Eco.Framework; исходники и логи этого проекта. Текст системного промпта хранится отдельно в DesignFiles/Prompt.','Caption')
OUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT)
copyfile(OUT,ROOT/'EcoLab2_Report.docx')
print(OUT)
